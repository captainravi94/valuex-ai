import io
import re
import requests
import openpyxl
import pandas as pd
import numpy as np

def _parse_screener_datasheet(wb, uploaded_file):
    ws = wb['Data Sheet']
    
    # 1. Company Name
    comp_name = ws.cell(1, 2).value
    if not comp_name:
        raw_name = uploaded_file.name.rsplit(".", 1)[0]
        comp_name = re.sub(r"[_\-\d]+", " ", raw_name).strip().title()
    else:
        comp_name = str(comp_name).strip().title()

    # 2. Extract Annual Date Columns (Row 16: P&L Report Date)
    date_cols = {}
    for col_idx in range(2, ws.max_column + 1):
        dt_val = ws.cell(16, col_idx).value
        if dt_val is not None:
            if hasattr(dt_val, "strftime"):
                label = dt_val.strftime("FY%y")
            else:
                label = str(dt_val).strip()
            date_cols[col_idx] = label

    if not date_cols:
        return None

    periods = list(date_cols.values())
    raw_data = {}

    def extract_row_series(row_idx):
        return {p: float(ws.cell(row_idx, c).value or 0.0) for c, p in date_cols.items()}

    # 3. Read Statements by Explicit Row Boundaries (Prevents Quarterly Overwrite)
    # P&L Section: Rows 17 to 35
    for r in range(17, 36):
        label = ws.cell(r, 1).value
        if label and str(label).strip():
            raw_data[str(label).strip().lower()] = extract_row_series(r)

    # Balance Sheet Section: Rows 57 to 75
    for r in range(57, 76):
        label = ws.cell(r, 1).value
        if label and str(label).strip():
            key = str(label).strip().lower()
            if key == "total":
                key = "total_assets" if "total_liabilities" in raw_data else "total_liabilities"
            raw_data[key] = extract_row_series(r)

    # Cash Flow Section: Rows 82 to 89
    for r in range(82, 90):
        label = ws.cell(r, 1).value
        if label and str(label).strip():
            raw_data[str(label).strip().lower()] = extract_row_series(r)

    # 4. Construct Canonical Statement Matrix
    canonical = {}

    # Revenue
    canonical["revenue"] = raw_data.get("sales", {p: 0.0 for p in periods})

    # EBITDA = PBT + Depreciation + Interest
    pbt = raw_data.get("profit before tax", {p: 0.0 for p in periods})
    dep = raw_data.get("depreciation", {p: 0.0 for p in periods})
    inte = raw_data.get("interest", {p: 0.0 for p in periods})
    other_inc = raw_data.get("other income", {p: 0.0 for p in periods})
    canonical["ebitda"] = {p: pbt[p] + dep[p] + inte[p] for p in periods}
    canonical["operating_ebitda"] = {p: max(canonical["ebitda"][p] - other_inc[p], 0.0) for p in periods}
    canonical["other_income"] = other_inc

    # PAT
    canonical["pat"] = raw_data.get("net profit", {p: 0.0 for p in periods})

    # Equity & Net Worth: Capital + Reserves (Resolves Net Worth Inconsistency)
    equity_cap = raw_data.get("equity share capital", {p: 0.0 for p in periods})
    reserves = raw_data.get("reserves", {p: 0.0 for p in periods})
    canonical["equity_share_capital"] = equity_cap
    canonical["reserves"] = reserves
    canonical["total_equity"] = {p: equity_cap[p] + reserves[p] for p in periods}

    # Borrowings & Debt
    canonical["total_borrowings"] = raw_data.get("borrowings", {p: 0.0 for p in periods})

    # Assets & Liabilities
    canonical["total_assets"] = raw_data.get("total_assets", raw_data.get("total", {p: 0.0 for p in periods}))
    canonical["total_liabilities"] = raw_data.get("total_liabilities", canonical["total_assets"])

    # Working Capital Line Items
    canonical["trade_receivables"] = raw_data.get("receivables", {p: 0.0 for p in periods})
    canonical["inventories"] = raw_data.get("inventory", {p: 0.0 for p in periods})
    canonical["cash_and_equivalents"] = raw_data.get("cash & bank", {p: 0.0 for p in periods})

    # Operating Cash Flow (CFO)
    canonical["cfo"] = raw_data.get("cash from operating activity", {p: 0.0 for p in periods})

    # Derived Real Capex: Net Block(t) - Net Block(t-1) + Depreciation(t) + CWIP(t) - CWIP(t-1)
    net_block = raw_data.get("net block", {p: 0.0 for p in periods})
    cwip = raw_data.get("capital work in progress", {p: 0.0 for p in periods})
    capex_series = {}
    for i, p in enumerate(periods):
        if i == 0:
            capex_series[p] = dep[p]
        else:
            prev_p = periods[i - 1]
            nb_delta = net_block[p] - net_block[prev_p]
            cwip_delta = cwip[p] - cwip[prev_p]
            computed_capex = nb_delta + dep[p] + cwip_delta
            capex_series[p] = max(round(computed_capex, 2), 0.0)

    canonical["capex"] = capex_series
    canonical["fcf"] = {p: canonical["cfo"][p] - canonical["capex"][p] for p in periods}

    df = pd.DataFrame(canonical).T
    df.index.name = "metric"
    df.columns = periods

    currency = "$" if any(w in comp_name.lower() for w in ["inc", "corp", "apple", "tesla", "citi"]) else "₹"
    scale = "Cr" if currency == "₹" else "M"

    return df, f"Successfully parsed {len(periods)} financial periods for {comp_name}", comp_name, "Corporate / Industrial", currency, scale, {"periods": periods}

def ingest_from_excel_or_csv(uploaded_file):
    try:
        fname = uploaded_file.name.lower()
        if fname.endswith((".xlsx", ".xls")):
            uploaded_file.seek(0)
            wb = openpyxl.load_workbook(uploaded_file, data_only=True)
            if "Data Sheet" in wb.sheetnames:
                result = _parse_screener_datasheet(wb, uploaded_file)
                if result is not None:
                    return result

        uploaded_file.seek(0)
        if fname.endswith(".csv"):
            df_raw = pd.read_csv(uploaded_file, header=None)
        else:
            xl = pd.ExcelFile(uploaded_file)
            sheet_to_use = xl.sheet_names[0]
            for s in xl.sheet_names:
                if any(k in s.lower() for k in ["profit", "consolidated", "standalone", "sheet1"]):
                    sheet_to_use = s
                    break
            df_raw = pd.read_excel(xl, sheet_name=sheet_to_use, header=None)

        header_row_idx = None
        for idx, row in df_raw.iloc[:50].iterrows():
            row_items = [str(v).strip() for v in row.values if pd.notna(v)]
            matches = [v for v in row_items if re.search(r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec|fy|\b20\d\d\b)", v, re.IGNORECASE)]
            if len(matches) >= 2:
                header_row_idx = idx
                break

        if header_row_idx is None:
            header_row_idx = 0

        raw_headers = list(df_raw.iloc[header_row_idx].values)
        headers = []
        for h in raw_headers:
            if hasattr(h, "strftime"):
                headers.append(h.strftime("%b %Y"))
            elif pd.isna(h):
                headers.append("")
            else:
                headers.append(str(h).replace("\n", " ").strip())

        df = df_raw.iloc[header_row_idx + 1:].copy()
        
        seen = {}
        deduped_headers = []
        for h in headers:
            if h not in seen:
                seen[h] = 0
                deduped_headers.append(h)
            else:
                seen[h] += 1
                deduped_headers.append(f"{h}_{seen[h]}")
        df.columns = deduped_headers

        first_col = df.columns[0]
        df = df.rename(columns={first_col: "metric"})
        df = df[df["metric"].notna()]
        df["metric"] = df["metric"].astype(str).str.strip()
        df = df[~df["metric"].str.lower().str.contains("screener|http|www|report|source|notes", regex=True)]
        df = df[df["metric"] != ""]

        valid_cols = ["metric"]
        for col in df.columns[1:]:
            s = df[col].iloc[:, 0] if isinstance(df[col], pd.DataFrame) else df[col]
            num_count = pd.to_numeric(
                s.astype(str).str.replace(",", "", regex=False).str.replace("%", "", regex=False).str.strip(),
                errors="coerce"
            ).notna().sum()
            if num_count >= 2:
                valid_cols.append(col)

        df = df[valid_cols]

        period_cols = [c for c in df.columns if c != "metric"]
        for col in period_cols:
            s = df[col].iloc[:, 0] if isinstance(df[col], pd.DataFrame) else df[col]
            clean_s = (
                s.astype(str)
                .str.replace(",", "", regex=False)
                .str.replace("₹", "", regex=False)
                .str.replace("$", "", regex=False)
                .str.replace("(", "-", regex=False)
                .str.replace(")", "", regex=False)
                .str.strip()
            )
            df[col] = pd.to_numeric(clean_s, errors="coerce").fillna(0.0)

        df = df.set_index("metric")
        df = df[~df.index.duplicated(keep="first")]

        raw_name = uploaded_file.name.rsplit(".", 1)[0]
        comp_name = re.sub(r"[_\-\d]+", " ", raw_name).strip().title()
        currency = "$" if any(w in comp_name.lower() for w in ["inc", "corp", "apple", "tesla", "citi"]) else "₹"
        scale = "Cr" if currency == "₹" else "M"

        return df, f"Successfully parsed {len(df.columns)} financial periods for {comp_name}", comp_name, "Corporate / Industrial", currency, scale, {"periods": list(df.columns)}

    except Exception as e:
        return None, f"Ingestion error: {str(e)}", None, None, "₹", "Cr", {}

def fetch_from_ir_link(url: str):
    try:
        if not url.startswith("http"):
            url = "https://" + url
        headers = {"User-Agent": "Mozilla/5.0"}
        resp = requests.get(url, headers=headers, timeout=10)
        tables = pd.read_html(io.StringIO(resp.text))
        if tables:
            best = max(tables, key=lambda t: t.shape[0] * t.shape[1])
            return best, f"Extracted table ({best.shape[0]} rows x {best.shape[1]} cols)."
        return None, "No HTML tables detected"
    except Exception as e:
        return None, str(e)